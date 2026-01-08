import unittest
import os
from fpdf import FPDF
from tasks import process_pdf_layout_aware
from docx import Document

class TestLayoutAwarePipeline(unittest.TestCase):
    def setUp(self):
        """Set up a self-contained, multi-page dummy PDF for testing."""
        self.test_pdf_path = "self_contained_test.pdf"
        pdf = FPDF()

        # Page 1
        pdf.add_page()
        pdf.set_font("Arial", 'B', size=24)
        pdf.cell(200, 20, txt="Document Title", ln=True, align='C')
        pdf.ln(20)
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, txt="This is the first paragraph on the first page. It contains some introductory text to test the paragraph detection.")

        # Page 2
        pdf.add_page()
        pdf.set_font("Arial", 'B', size=18)
        pdf.cell(200, 15, txt="Chapter 1: A New Beginning", ln=True, align='L')
        pdf.ln(10)
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, txt="This is the main content of the second page, which follows a chapter heading. The layout model should be able to distinguish this from the title on the first page.")

        pdf.output(self.test_pdf_path)
        os.makedirs("results", exist_ok=True)

    def tearDown(self):
        """Clean up the created files."""
        if os.path.exists(self.test_pdf_path):
            os.remove(self.test_pdf_path)

        docx_path = f"results/{os.path.splitext(os.path.basename(self.test_pdf_path))[0]}_structured.docx"
        if os.path.exists(docx_path):
            os.remove(docx_path)

    def test_pipeline_with_generated_pdf(self):
        """
        Test the full pipeline with a programmatically generated, multi-page PDF.
        """
        result = process_pdf_layout_aware(self.test_pdf_path, "self_contained_test.pdf")

        # 1. Check status and file existence
        self.assertEqual(result['status'], 'SUCCESS', f"Processing failed: {result.get('message')}")
        self.assertIsNotNone(result['docx_path'])
        docx_path = result['docx_path']
        self.assertTrue(os.path.exists(docx_path))

        # 2. Verify DOCX content
        doc = Document(docx_path)
        # We expect at least 3 blocks from our generated 2-page PDF
        self.assertTrue(len(doc.paragraphs) >= 3, f"Expected >= 3 paragraphs, but found {len(doc.paragraphs)}")

        doc_text = "\n".join([p.text for p in doc.paragraphs])
        self.assertIn("Document Title", doc_text)
        self.assertIn("A New Beginning", doc_text)
        self.assertIn("introductory text", doc_text)

if __name__ == '__main__':
    unittest.main()
